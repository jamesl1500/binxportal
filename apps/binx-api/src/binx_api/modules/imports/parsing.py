"""
Turns an uploaded .csv / .xlsx file into a header row plus data rows of
plain strings — nothing else. Field mapping and validation happen later
(see service.py), against JSON the client sends back, so this module never
touches the database and never stores the file.

Defensive by design, since the bytes come straight from a user:
- a hard size cap before anything is parsed,
- .xlsx is read with openpyxl in read-only, values-only mode (formulas are
  never evaluated, macros never loaded), after checking the zip's total
  uncompressed size so a tiny "zip bomb" can't balloon in memory,
- row/column caps so a huge sheet fails fast with a clear message.
"""

from __future__ import annotations

import csv
import io
import zipfile
from dataclasses import dataclass
from datetime import date, datetime, time

MAX_FILE_BYTES = 2 * 1024 * 1024
MAX_XLSX_UNCOMPRESSED_BYTES = 50 * 1024 * 1024
MAX_COLUMNS = 50
MAX_CELL_CHARS = 10_000


class ImportFileError(ValueError):
    """The upload can't be read as a spreadsheet. ``status_code`` is what the
    router should answer with; the message is shown to the user verbatim."""

    def __init__(self, message: str, status_code: int = 400) -> None:
        super().__init__(message)
        self.status_code = status_code


@dataclass
class ParsedSheet:
    headers: list[str]
    # (spreadsheet row number, cells) — the row number is what users see in
    # Excel, so error messages can point at "row 14" meaningfully.
    rows: list[tuple[int, list[str]]]


def parse_upload(file_name: str, content: bytes, *, max_rows: int) -> ParsedSheet:
    if len(content) > MAX_FILE_BYTES:
        raise ImportFileError(f"Files must be {MAX_FILE_BYTES // (1024 * 1024)}MB or smaller", 413)
    if not content:
        raise ImportFileError("That file is empty")

    lower = file_name.lower()
    if lower.endswith(".xlsx"):
        raw_rows = _read_xlsx(content, max_rows=max_rows)
    elif lower.endswith(".csv") or lower.endswith(".txt"):
        raw_rows = _read_csv(content, max_rows=max_rows)
    elif lower.endswith(".xls"):
        raise ImportFileError("Old .xls files aren't supported — save it as .xlsx or .csv and try again", 415)
    else:
        raise ImportFileError("Upload a .csv or .xlsx file", 415)

    return _to_sheet(raw_rows, max_rows=max_rows)


def _cell_to_str(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, float) and value.is_integer():
        # Excel stores phone numbers / IDs typed as numbers as floats.
        return str(int(value))
    if isinstance(value, datetime):
        return value.date().isoformat() if value.time() == time(0) else value.isoformat(sep=" ")
    if isinstance(value, date):
        return value.isoformat()
    return str(value)


def _read_xlsx(content: bytes, *, max_rows: int) -> list[list[str]]:
    # openpyxl is imported lazily: it's only needed for this one code path.
    from openpyxl import load_workbook

    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            if sum(info.file_size for info in archive.infolist()) > MAX_XLSX_UNCOMPRESSED_BYTES:
                raise ImportFileError("That spreadsheet is too large to import", 413)
    except zipfile.BadZipFile:
        raise ImportFileError("That file isn't a valid .xlsx spreadsheet") from None

    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except ImportFileError:
        raise
    except Exception:  # openpyxl raises a zoo of types for malformed files
        raise ImportFileError("That file isn't a valid .xlsx spreadsheet") from None

    try:
        sheet = workbook.worksheets[0] if workbook.worksheets else None
        if sheet is None:
            raise ImportFileError("That spreadsheet has no sheets")
        rows: list[list[str]] = []
        filled = 0
        for values in sheet.iter_rows(values_only=True):
            cells = [_cell_to_str(value) for value in values]
            rows.append(cells)
            filled += any(cell.strip() for cell in cells)
            if filled > max_rows + 1:  # header + one past the cap: enough to report "too many"
                break
        return rows
    finally:
        workbook.close()


def _decode(content: bytes) -> str:
    try:
        return content.decode("utf-8-sig")
    except UnicodeDecodeError:
        # Excel's "CSV" export on Windows is cp1252, not UTF-8.
        return content.decode("cp1252", errors="replace")


def _read_csv(content: bytes, *, max_rows: int) -> list[list[str]]:
    text = _decode(content)
    try:
        dialect: type[csv.Dialect] | csv.Dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    rows: list[list[str]] = []
    filled = 0
    try:
        for cells in csv.reader(io.StringIO(text), dialect):
            rows.append(cells)
            filled += any(cell.strip() for cell in cells)
            if filled > max_rows + 1:  # header + one past the cap: enough to report "too many"
                break
    except csv.Error:
        raise ImportFileError("That CSV file couldn't be read — check it isn't corrupted") from None
    return rows


def _to_sheet(raw_rows: list[list[str]], *, max_rows: int) -> ParsedSheet:
    numbered = [
        (index + 1, [cell.strip()[:MAX_CELL_CHARS] for cell in cells])
        for index, cells in enumerate(raw_rows)
        if any(cell.strip() for cell in cells)
    ]
    if not numbered:
        raise ImportFileError("That file has no rows")

    _header_row, headers = numbered[0]
    # Trailing empty header cells are just formatting leftovers.
    while headers and not headers[-1]:
        headers.pop()
    if not headers:
        raise ImportFileError("The first row should be column headings (e.g. “Name”, “Email”)")
    if len(headers) > MAX_COLUMNS:
        raise ImportFileError(f"That file has more than {MAX_COLUMNS} columns")

    data = numbered[1:]
    if not data:
        raise ImportFileError("That file only has a header row — add at least one row below it")
    if len(data) > max_rows:
        raise ImportFileError(f"You can import up to {max_rows} rows at a time — split the file and try again", 413)

    width = len(headers)
    rows = [(row_number, (cells + [""] * width)[:width]) for row_number, cells in data]
    return ParsedSheet(headers=[header or f"Column {index + 1}" for index, header in enumerate(headers)], rows=rows)
